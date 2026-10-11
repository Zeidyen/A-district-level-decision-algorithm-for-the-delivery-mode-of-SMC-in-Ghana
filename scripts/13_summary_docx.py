"""Results summary: every number computed from data/derived, written once as Markdown
(../results_summary.md) and rendered to Word (../results_summary.docx). The Claude Doc
version is filled from the same Markdown, section by section.

Run after the pipeline (run_all.sh) and the scenario runs (10_step4_impact.R <scenario>,
14_sensitivity.R, robustness table, 19/20 DHIMS scripts, 11/12 figures).
"""
import re
import pandas as pd
from scipy.stats import spearmanr

D = "data/derived/"
OUT_MD, OUT_DOCX = "../results_summary.md", "../results_summary.docx"
AS_OF = "10 October 2026"
NORTH = ["Northern", "Savannah", "North East", "Upper East", "Upper West"]

# ---- numbers (nothing typed; literature constants are in the k table and sources) -----
rb = pd.read_csv(D + "district_robustness.csv")
s4 = pd.read_csv(D + "district_step4.csv")
sens = pd.read_csv(D + "step4_sensitivity.csv")
stress = pd.read_csv(D + "step4_k_stress.csv").set_index("k")
be = pd.read_csv(D + "step4_breakeven.csv").set_index("region")["breakeven_k"]
rbe = pd.read_csv(D + "region_breakeven.csv").set_index("region")           # 21_breakeven_gsa.R
dbe = pd.read_csv(D + "district_breakeven.csv")
favc = pd.read_csv(D + "favourable_check.csv").set_index("setting")
rcc = pd.read_csv(D + "region_breakeven_cc.csv").set_index("region")
dcv = pd.read_csv(D + "district_cv_summary.csv").set_index("model")      # 22_district_cv.R
mls = pd.read_csv(D + "ml_stack_cv_summary.csv").set_index("model")      # 26_ml_stack_cv.R
slcv = pd.read_csv(D + "slip_ml_cv.csv")                                # 27_slip_drivers.py
slg = pd.read_csv(D + "slip_shap_groups.csv").set_index("group")["share_north"]
slr = pd.read_csv(D + "slip_shap_by_region.csv")
s4s = pd.read_csv(D + "district_step4_stacked.csv")
n_contacts = len(pd.read_csv(D + "slip_contacts_features.csv"))
sso = pd.read_csv(D + "season_slip_or.csv").set_index("model")              # 28_season_slips.R
ssm = pd.read_csv(D + "season_slip_months.csv").set_index("month")
ssf = pd.read_csv(D + "season_slip_factor.csv").iloc[0]
hyb = pd.read_csv(D + "hybrid_summary.csv").set_index("set")               # 29_hybrid.R
hyb_first = pd.read_csv(D + "hybrid_curve.csv")
eqg = pd.read_csv(D + "equity_by_group.csv")                               # 30_equity.py
eqs = pd.read_csv(D + "equity_summary.csv").set_index("k")
alg = pd.read_csv(D + "delivery_algorithm.csv")                            # 31_delivery_algorithm.R
algc = pd.read_csv(D + "delivery_algorithm_counts.csv")
gsa = pd.read_csv(D + "gsa_indices.csv")
gsr = pd.read_csv(D + "gsa_region_ratio.csv").set_index("region")
fx = pd.read_csv(D + "step2_fixed.csv").set_index("term")
cv = pd.read_csv(D + "cv_results.csv")
rain = pd.read_csv(D + "rain_or.csv").set_index("term")
age = pd.read_csv(D + "age_bands_vitA.csv").set_index("band")
ager = pd.read_csv(D + "age_reach.csv").iloc[0]
rbu = pd.read_csv(D + "region_burden.csv").set_index("region")
fp = pd.read_csv(D + "fingerprint_vs_map.csv")
val = pd.read_csv(D + "fingerprint_validation.csv").set_index("test")
u5pp = pd.read_csv(D + "dhims_u5_season_prepost.csv")
u5n = u5pp[u5pp.Region.isin(["Northern", "Savannah", "North East"])]
u5ue = u5pp[u5pp.Region == "Upper East"]
u5_cls = u5n.merge(fp[["key", "class"]], left_on="adm2_key", right_on="key").groupby("class")["change"].median()
wet = pd.read_csv(D + "wet_access_summary.csv")
acc = pd.read_csv(D + "district_access_u5.csv")
fac = pd.read_csv(D + "facilities.csv")
fac_n = fac[fac.region.isin(["Northern", "Savannah", "North East", "Upper East", "Upper West"])]
def near(measure, band="<30 min"):
    a = acc[acc.measure == measure].copy(); a["reg"] = a.region.str.replace(" Region", "", regex=False)
    t = a.groupby("reg").u5.sum(); b = a[a.band == band].groupby("reg").u5.sum()
    return (b / t), b.sum() / t.sum()
chps30_reg, chps30_all = near("tt_walk_chps")
cl = pd.read_csv(D + "dhs_clusters.csv").merge(pd.read_csv(D + "dhs_clusters_timeliness.csv"), on="cluster")

n_var = int(rb["n"].iloc[0]); n_d = len(rb); kids = rb["kids"].sum()
pct = lambda x: f"{100 * x:.0f}%"
k1000 = lambda x: f"{x / 1000:,.0f},000"
grp = rb.apply(lambda r: "all" if r.red == r.n else ("red9" if r.red >= 9 else ("never" if r.green == 0 else "mixed")), axis=1)
rb["grp"] = grp
def names(g): return ", ".join(sorted(rb.loc[rb.grp == g, "district"].str.replace(r"\s+", " ", regex=True)))
def kidsg(g): return rb.loc[rb.grp == g, "kids"].sum()
never = rb.green.eq(0); n_never = int(never.sum()); pct_never = rb.loc[never, "kids"].sum() / kids
mixed_regions = sorted(rb.loc[rb.grp == "mixed", "region"].unique())
# Describe the "could switch" group and the never-green group from the data, not by assumption.
inside = rb.region.isin(mixed_regions)
n_out_never = int((~inside & never).sum()); n_out = int((~inside).sum())
in_never = sorted(rb.loc[inside & never, "district"].str.replace(r"\s+", " ", regex=True))
not_mixed_in = sorted(rb.loc[inside & (rb.grp != "mixed"), "district"].str.replace(r"\s+", " ", regex=True))
mixed_desc = (f"every district of {' and '.join(mixed_regions)}" if not not_mixed_in else
              f"every district of {' and '.join(mixed_regions)} except {' and '.join([', '.join(not_mixed_in[:-1]), not_mixed_in[-1]] if len(not_mixed_in) > 1 else not_mixed_in)}")
outside_clause = (f"all {n_out} districts outside {' and '.join(mixed_regions)}" if n_out_never == n_out else
                  f"{n_out_never} of the {n_out} districts outside {' and '.join(mixed_regions)}")
inside_extra = f", plus {' and '.join(in_never)}" if in_never else ""
never_sentence = ((f"none of the {n_out} districts there is ever classed green." if n_out_never == n_out else
                   f"{n_out_never} of the {n_out} districts there are never classed green.") +
                  (f" With {' and '.join(in_never)}, these {n_never} never-green districts hold {pct(pct_never)} of children aged 3–59 months."
                   if in_never else f" They hold {pct(pct_never)} of children aged 3–59 months."))

main = sens.set_index("scenario").loc["main"]
fm = pd.read_csv(D + "district_step4_facmaina.csv")
_fc = fm[["district", "class"]].merge(s4[["district", "class"]], on="district", suffixes=("_old", "_new"))
n_fac_changed = int((_fc.class_old != _fc.class_new).sum())
cls = s4["class"].value_counts()
rcov = s4.groupby("region").apply(lambda x: (x.routine_cov * x.children_3_59m).sum() / x.children_3_59m.sum())
allcov = (s4.routine_cov * s4.children_3_59m).sum() / s4.children_3_59m.sum()

nc = cl[cl.region.isin(NORTH) & (cl.contacts > 0)].copy()
nc["band"] = pd.cut(nc.tt_walk_any, [0, 30, 60, 120, 1e9], right=False, labels=["<30", "30-60", "60-120", ">120"])
qb = nc.groupby("band", observed=True).apply(lambda x: x.slip56.sum() / x.contacts.sum())
q_north = nc.slip56.sum() / nc.contacts.sum()

rho_all = spearmanr(fp["index"], fp["gap_slip"]); rho_ratio = spearmanr(fp["index"], fp["ratio"])
res = lambda c: fp[c] - fp.groupby("Region")[c].transform("mean")
rho_within = spearmanr(res("index"), res("gap_slip"))
fp_cls = fp.groupby("class")["index"].median()
fp_reg = fp.groupby("Region")["index"].median()
fp_weak = fp.sort_values("index", ascending=False)

tt = fx.loc["travel_time_doubling"]
cvn = cv[(cv.subset == "north")].set_index("model")
w_all = wet[(wet["mode"] == "walk") & (wet.region == "All north")].set_index("rule")
adh_r = 0.845 * 0.8674 / 0.9850
share_old = float(re.search(r"aged 24-59 m ([0-9.]+)", open("out/logs/10_step4_age.log").read()).group(1))

# Break-even k (best case for routine delivery: no age gap, no adherence loss).
kidsb = dbe.children_3_59m.sum()
b_lt1 = dbe.be_best < 1; b_low = (dbe.be_best >= 1) & (dbe.be_best < 1.4)
b_ev = (dbe.be_best >= 1.4) & (dbe.be_best < 2.8); b_hi = dbe.be_best >= 2.8
bpk = lambda m: pct(dbe.loc[m, "children_3_59m"].sum() / kidsb)
andj = lambda xs: ", ".join(xs[:-1]) + " and " + xs[-1] if len(xs) > 1 else "".join(xs)
b_ev_regions = sorted(dbe.loc[b_ev, "region"].unique())
b_hi_names = ", ".join(dbe.loc[b_hi, "district"].str.replace(r"\s+", " ", regex=True))
n_real_lt1 = int((dbe.be_real < 1).sum()); n_real0 = int((dbe.be_real == 0).sum())
fav1 = favc.loc["k=1, all else favourable, C_c mean"]
fav14 = favc.loc["k=1.4 (evidence lower bound), favourable, C_c mean"]
# Break-even campaign coverage (21_breakeven_gsa.R): routine reaches 90% of the campaign only if
# the campaign's per-cycle coverage is below this.
cc_rows = "\n".join(
    f"| {r} | {pct(rcc.loc[r, 'cc_k14_best'])} | {pct(rcc.loc[r, 'cc_k2_best'])} "
    f"({pct(rcc.loc[r, 'cc_k2_best_lo'])}–{pct(rcc.loc[r, 'cc_k2_best_hi'])}) | {pct(rcc.loc[r, 'cc_k2_real'])} |"
    for r in NORTH)
cc75 = dbe.cc_k2_best < 0.75; cc87 = dbe.cc_k2_best < 0.87
be_rows = "\n".join(f"| {r} | {rbe.loc[r, 'best']:.2f} | {rbe.loc[r, 'best_red']:.2f} |" for r in NORTH)

# Global sensitivity: first-order variance shares by region.
G = gsa.pivot(index="factor", columns="region", values="first_order")[NORTH]
FK, FA, FC, FQ = ("k (routine vs vaccination slip)", "age reach (24-59 m)", "campaign coverage",
                  "step-2 timeliness surface")
short = {FK: "k", FA: "age reach", FC: "campaign coverage", FQ: "the timeliness map",
         "first-contact reach": "first-contact reach"}
small = G.loc[["routine adherence", "partial-course protection"]]
top = {r: G[r].idxmax() for r in NORTH}
k_top = [r for r in NORTH if top[r] == FK]; not_k_top = [r for r in NORTH if top[r] != FK]
not_k_text = "".join(f" In {r}, {short[top[r]]} matters more ({pct(G.loc[top[r], r])}; k {pct(G.loc[FK, r])})."
                     for r in not_k_top)
gsa_rows = "\n".join(
    f"| {r} | {gsr.loc[r, 'ratio_median']:.2f} ({gsr.loc[r, 'lo']:.2f}–{gsr.loc[r, 'hi']:.2f}) | "
    f"{pct(gsr.loc[r, 'p_green'])} | {pct(gsr.loc[r, 'p_red'])} | {short[top[r]]} ({pct(G.loc[top[r], r])}) |"
    for r in NORTH)

# Machine learning (24-27).
_st = s4[["district", "class"]].merge(s4s[["district", "class", "ratio"]], on="district", suffixes=("", "_stk"))
n_stk_same = int((_st["class"] == _st.class_stk).sum())
stk_cls = s4s["class"].value_counts()
stk_moves = "; ".join(f"{r.district.strip()} {r['class']} to {r.class_stk}" for _, r in _st[_st["class"] != _st.class_stk].iterrows())
ratio_corr = s4[["district", "ratio"]].merge(s4s[["district", "ratio"]], on="district").filter(like="ratio").corr().iloc[0, 1]
auc = lambda m: slcv[(slcv.model == m) & (slcv.subset == "north")].auc.iloc[0]
eff = lambda f, r: slr[(slr.feature == f) & (slr.region == r)].effect.iloc[0]
ml_claim = ("**Machine learning predicts district slip rates better than the spatial model**"
            if mls.loc["ml_alone_stack", "wilcoxon_p_vs_main"] < 0.05 else
            "**Machine learning gives lower prediction error than the spatial model, though the district-by-district difference is not statistically conclusive**")
mlrows = "\n".join(f"| {lab} | {mls.loc[m, 'mae']:.3f} | {mls.loc[m, 'corr']:.2f} | {mls.loc[m, 'logscore']:.3f} |"
                   for m, lab in [("ml_alone_stack", "Machine-learning stack alone"), ("stacked", "Machine-learning stack + spatial field"),
                                  ("main", "Main spatial model"), ("region", "Regional means")])

# Season of contact, hybrid delivery, equity (28-30).
so = sso.loc["north, adjusted"]
hg = hyb.loc["green (ratio >= 0.9)"]; h2 = hyb.loc["loss <= 2%"]; h5 = hyb.loc["loss <= 5%"]
hyb_rows = "\n".join(f"| {lab} | {int(hyb.loc[k, 'districts'])} | {hyb.loc[k, 'pct_children']:.0f}% | {hyb.loc[k, 'pct_loss']:.1f}% | {hyb.loc[k, 'pct_loss_age_adh']:.1f}% | {hyb.loc[k, 'regions']} |"
                     for k, lab in [("green (ratio >= 0.9)", "Green districts only"), ("loss <= 1% of cases averted", "Up to 1% of cases averted lost"),
                                    ("loss <= 2%", "Up to 2% lost"), ("loss <= 5%", "Up to 5% lost"), ("all districts", "All districts")])
eq2 = eqg[(eqg.k == 2.0) & (eqg.by == "Wealth quintile")].set_index("group"); eq2.index = eq2.index.astype(int)
eqe = eqg[(eqg.k == 2.0) & (eqg.by == "Mother's education")].set_index("group")
QL = {1: "Poorest", 2: "Poorer", 3: "Middle", 4: "Richer", 5: "Richest"}
eq_rows = "\n".join(f"| {QL[g]} | {pct(eq2.loc[g, 'share_of_children'])} | {eq2.loc[g, 'routine_cov']:.2f} | {eq2.loc[g, 'lost_per_1000']:.0f} | {pct(eq2.loc[g, 'share_of_loss'])} |" for g in range(1, 6))
ao = alg[alg.model == "outreach"]
ac = algc[algc.model == "outreach"].set_index("outcome")
def acn(o): return f"{int(ac.loc[o, 'districts'])} districts ({ac.loc[o, 'pct_children']:.0f}% of children)"
pilot_top = ", ".join(ao[ao.outcome == "Pilot"].sort_values("pilot_priority").district.head(6))
meas_first = ", ".join(ao[ao.outcome == "Measure first"].district)
flags = "; ".join(f"{r.district} ({r.robustness_flag.replace('check: ', '')})" for _, r in ao[ao.robustness_flag.fillna("") != ""].iterrows())
meas_by_reg = alg[(alg.model == "outreach") & (alg.measure_first.fillna("") != "")].drop_duplicates("region").set_index("region").measure_first
n_eq = int(eqg[(eqg.k == 2.0) & (eqg.by == "Wealth quintile")].children.sum())

def moved(df):
    """'fell from 43% to 41%' / 'rose from ...' / 'did not change (37%)' from the district medians."""
    a, b = df.share_pre.median(), df.share_post.median()
    if round(100 * a) == round(100 * b):
        return f"did not change ({pct(a)})"
    return f"{'fell' if b < a else 'rose'} from {pct(a)} to {pct(b)}"

def ci(o, l, h, d=2): return f"{o:.{d}f} ({l:.{d}f}–{h:.{d}f})"

# ---- text ----------------------------------------------------------------------------
S = []  # (key, markdown) per section; the doc and the Word file are built from these

S.append(("lead", f"""Outside {' and '.join(mixed_regions)}, routine delivery of SMC through health facilities and CHPS cannot replace the door-to-door campaign under any of the {n_var} model variants tested: {never_sentence} In {int(b_lt1.sum())} districts holding {bpk(b_lt1)} of children, routine delivery would match the campaign only if a routine SMC visit were missed *less* often than a vaccination visit, which no evidence suggests; this holds whatever the true value of the uncertain reliability factor k. Inside {' and '.join(mixed_regions)}, routine delivery could come close to the campaign only if routine monthly contacts become nearly as reliable as vaccination visits and reach children aged 2–4. Everywhere, routine delivery would shift protection away from the poorest children: children in Ghana's poorest wealth quintile would bear {pct(eq2.loc[1, 'share_of_loss'])} of the protection lost."""))

S.append(("found", f"""## What we found

Each district is classed by the cases averted under routine delivery as a share of those the campaign averts: **green** 90% or more (could switch), **amber** 70–90%, **red** below 70% (keep the campaign). Because the green count swings from {sens.green.min()} to {sens.green.max()} across the {n_var} variants, the result is reported as how consistently each district is classed, not as a single count.

| Group | Districts | Children 3–59 m | Which |
| --- | --- | --- | --- |
| Red in all {n_var} variants | {(rb.grp == 'all').sum()} | {k1000(kidsg('all'))} | {names('all')} |
| Red in at least 9 of {n_var} | {(rb.grp == 'red9').sum()} | {k1000(kidsg('red9'))} | {names('red9')} |
| Never green | {(rb.grp == 'never').sum()} | {k1000(kidsg('never'))} | {names('never')} |
| Could switch under some assumptions | {(rb.grp == 'mixed').sum()} | {k1000(kidsg('mixed'))} | {mixed_desc} |

In the main analysis {cls.get('green', 0)} districts are green, {cls.get('amber', 0)} amber and {cls.get('red', 0)} red; only {int(main.green_p80)} is green with probability 0.8 or more, against {int(main.red_p80)} red. Switching every district to routine delivery would avert about {main.shortfall_pct:.0f}% fewer cases than the campaign."""))

S.append(("map", f"""## Map

The divide runs between the Upper East and Upper West plateau and the low Volta basin to the south and east. The deepest red is eastern Northern Region.

![District class across {n_var} model variants](out/summary_district_map.png)

Large districts dominate the map visually but hold few children; the table above gives the counts that matter for planning."""))

S.append(("why", f"""## Why routine delivery falls short

Three things, each measured separately, push in the same direction.

- **Late monthly contacts, even close to a CHPS.** Card-dated Penta doses (DHS 2022) show {pct(q_north)} of monthly routine contacts in the north slipping by a month or more or never happening. Most children are close to a routine platform: {pct(chps30_all)} of under-5s in the five regions live within 30 minutes' walk of a CHPS (from {pct(chps30_reg.min())} in {chps30_reg.idxmin()} to {pct(chps30_reg.max())} in {chps30_reg.idxmax()}). Yet slips occur at every distance, rising only from {pct(qb['<30'])} within 30 minutes of a facility to {pct(qb['>120'])} beyond 2 hours. The barrier is how reliably the service delivers each month, more than distance. Contacts due in the SMC season slip no more often than at other times of year (adjusted odds ratio {so['or']:.2f}).
- **Routine services barely reach children aged 2–4.** Routine immunisation ends at 18 months, yet {pct(share_old)} of SMC-eligible children are 24–59 months. Vitamin A in the last 6 months, the nearest routine contact for them, falls from {age.loc['12-23','vitA']:.0f}% at 12–23 months to {age.loc['48-59','vitA']:.0f}% at 48–59 months (ratio {ager.ratio:.2f}, 95% CI {ager.lo:.2f}–{ager.hi:.2f}).
- **Lower adherence when doses go home.** With no follow-up visit, day-2 and day-3 adherence is expected to fall from about 84.5% (Upper East campaign) to about {pct(adh_r)}, based on the Nigerian gap between door-to-door and other delivery.

The model scales the vaccination slip rate by k, how much more often a routine malaria-product visit is missed than a vaccination visit, set at a median of 2 from Ghana's malaria vaccine rollout. Adding the age gap removes every green district; adding age and adherence together makes {int(sens.set_index('scenario').loc['age_adh','red'])} of {n_d} districts red."""))

S.append(("breakeven", f"""## How reliable would routine SMC have to be?

The main uncertainty is k: how much more often a routine SMC visit would be missed than a vaccination visit. The evidence puts it near 2 (1.4–2.8), but no study measures it for SMC. Turning the question round avoids having to know it: for each district, how large can k be before routine delivery falls below 90% of the campaign? The figures below take the best case for routine delivery (no age gap, no adherence loss).

| Region | Break-even k, 90% line | Break-even k, 70% line |
| --- | --- | --- |
{be_rows}

- **{int(b_lt1.sum())} districts ({bpk(b_lt1)} of children) have a break-even below 1.** Routine delivery would reach 90% of the campaign there only if routine SMC were more reliable than routine vaccination.
- **{int(b_low.sum())} more ({bpk(b_low)}) lie between 1 and 1.4**, below the lowest evidence for k.
- **{int(b_ev.sum())} ({bpk(b_ev)}) lie within the evidence range**, all in {andj(b_ev_regions)}; {int(b_hi.sum())} more ({b_hi_names}) lies above it. These are the only districts where the true value of k decides the answer.
- **Adding the age gap and the adherence loss puts every district below 1** ({n_real_lt1} of {n_d}); in {n_real0}, routine delivery cannot reach 90% even if no routine contact is ever late. Routine delivery that does not reach children aged 2–4 cannot match the campaign anywhere.
- **Even with everything set in routine delivery's favour** (k = 1, no age gap, no adherence loss), no district is red, but {fav1.pct_children_not_green:.0f}% of children live in districts that are still not green. At k = 1.4, the lower end of the evidence, {int(fav14.red)} districts are red and {fav14.pct_children_not_green:.0f}% of children are in districts that are not green.

![Break-even k by district, best case for routine delivery](out/breakeven_map.png)"""))

S.append(("campaign", f"""## How weak would the campaign have to be?

Routine delivery is judged against the campaign, so the answer also depends on how well the campaign itself performs. Only Upper East's campaign coverage has been measured (87% per cycle); elsewhere 87% is assumed. Routine delivery reaches 90% of the campaign only where the campaign covers fewer children per cycle than:

| Region | k = 1.4, best case | k = 2, best case (95% range) | k = 2, with age gap and adherence |
| --- | --- | --- | --- |
{cc_rows}

- **In Northern, the campaign would have to cover under {pct(rcc.loc['Northern', 'cc_k2_best'])} of children per cycle** at the central k, and under {pct(rcc.loc['Northern', 'cc_k14_best'])} even at the lowest evidence for k, far below the 87% measured in Upper East.
- **At k = 2, {int(cc75.sum())} districts ({bpk(cc75)} of children) would need a campaign below 75%**, the low end of the range assumed outside Upper East; {int(cc87.sum())} ({bpk(cc87)}) would need it below 87%.
- **In Upper East and Upper West the margin is thin:** {pct(rcc.loc['Upper East', 'cc_k2_best'])} and {pct(rcc.loc['Upper West', 'cc_k2_best'])} at k = 2, close to the measured 87%, so small differences in campaign performance decide the answer there.
- **The conclusion therefore assumes that campaigns outside Upper East perform roughly as well as the one measured.** If NMEP's district data show much lower coverage, the problem is the quality of the campaign rather than the choice of delivery mode.

![Break-even campaign coverage by region](out/breakeven_cc.png)"""))

S.append(("drivers", f"""## What drives the answer

All the uncertain inputs were varied together (k, campaign coverage, age reach, routine adherence, partial-course protection, first-contact reach and the timeliness map; 20,000 draws), and the variance in each region's routine-to-campaign ratio was split by input.

| Region | Ratio, median (95% range) | Chance green | Chance red | Largest driver |
| --- | --- | --- | --- | --- |
{gsa_rows}

- **k is the largest driver in {andj(k_top)}** ({pct(G.loc[FK, k_top].min())}–{pct(G.loc[FK, k_top].max())} of the variance).{not_k_text} Age reach explains {pct(G.loc[FA].min())}–{pct(G.loc[FA].max())} across regions.
- **The timeliness map explains {pct(G.loc[FQ].min())}–{pct(G.loc[FQ].max())}**: the geospatial model matters but does not decide the answer.
- **Adherence and partial-course protection explain {pct(small.values.min())}–{pct(small.values.max())} each**, and drug protection cancels from the ratio altogether.
- **What to measure first** follows from these shares; the algorithm below sets it out region by region.

![Share of variance in the routine-to-campaign ratio, by input and region](out/gsa_chart.png)"""))

S.append(("algorithm", f"""## The SMC delivery-mode algorithm

The analyses combine into a district-level decision rule that a national malaria programme can run inside its annual planning cycle and update each year. For each SMC-eligible district:

1. **Campaign check.** If campaign coverage per cycle, measured by household survey, is below the programme target (80%), fix the campaign first. Administrative coverage is usually 95–100% or more and cannot tell a good campaign from a weak one, so it is not used as measured coverage. The 80% target is only this quality gate; step 3 compares routine delivery with the measured campaign coverage where it exists (Upper East, 87%) and with an assumed 87% (75–95%) elsewhere, so no district is sent back at step 1.
2. **Age reach.** If the routine model cannot reach children aged 2–4 (clinic or immunisation contacts), keep the campaign.
3. **Reliability.** Compare the district's break-even k, with its uncertainty, with the evidence for k (1.4–2.8): keep the campaign if it is at least 80% likely to lie below 1.4; consider a switch if at least 80% likely to lie above 2.8; measure first if it probably lies below 1.4 but less surely; otherwise pilot.
4. **Equity safeguard.** If the poorest wealth quintile would bear significantly more than its share of the lost protection (the lower 95% bound of its loss share divided by its share of children is above 1; no WHO or published threshold exists for this ratio), a switch becomes a hybrid (routine delivery plus a campaign mop-up for the poorest and most remote).
5. **Pilot priority.** Rank pilot districts by their routine-to-campaign ratio.
6. **Measure and decide.** Measure the region's main drivers first; switch only if the measured k is below the district's break-even and age reach and adherence are confirmed.

| Outcome | Routine model reaching children aged 2–4 | Clinic or immunisation-linked model |
| --- | --- | --- |
| Keep campaign | {acn('Keep campaign')} | all {int(algc[(algc.model == 'clinic') & (algc.outcome == 'Keep campaign')].districts.iloc[0])} districts |
| Measure first | {acn('Measure first')} | none |
| Pilot | {acn('Pilot')} | none |
| Hybrid or switch | {int(ac.loc['Hybrid', 'districts'] + ac.loc['Switch', 'districts'])} districts | none |

- **Pilot first in** {pilot_top}.
- **Measure first in** {meas_first}: probably keep the campaign, but the answer rests on campaign coverage that has not been measured outside Upper East. They resolve to keep the campaign if survey coverage, or administrative coverage corrected against a survey, reaches about {pct(ao[ao.outcome == 'Measure first'].cc_threshold.min())}–{pct(ao[ao.outcome == 'Measure first'].cc_threshold.max())} per cycle. Inflated administrative coverage can only make routine delivery look worse, never better.
- **What to measure first** (in order): in Upper West, {meas_by_reg.get('Upper West', '').replace('; ', ', then ').replace('2-4', '2–4')}; in Upper East, {meas_by_reg.get('Upper East', '').replace('; ', ', then ').replace('2-4', '2–4')}.
- **Robustness flags** (outcome at odds with the 14 model variants): {flags}.
- **Every region triggers the equity safeguard**: the poorest quintile's share of the loss is {eqg[(eqg.k == 2.0) & (eqg.by == 'Region: poorest quintile')].eval('share_of_loss / share_of_children').min():.1f}–{eqg[(eqg.k == 2.0) & (eqg.by == 'Region: poorest quintile')].eval('share_of_loss / share_of_children').max():.1f} times its share of children, so any pilot should test a safeguard.

![The SMC delivery-mode algorithm, with northern Ghana counts](out/delivery_algorithm_flow.png)

![Algorithm outcome by district, under the two routine models](out/delivery_algorithm_map.png)"""))

S.append(("hybrid", f"""## A partial switch: hybrid delivery

Districts were moved from the campaign to routine delivery one at a time, best routine-to-campaign ratio first, to see how many children could move for how little protection.

| Move | Districts | Children moved | Cases averted lost | Lost, with age gap and adherence | Regions |
| --- | --- | --- | --- | --- | --- |
{hyb_rows}

- **Moving the {int(hg.districts)} green districts ({hg.pct_children:.0f}% of children) costs {hg.pct_loss:.1f}% of the campaign's cases averted**, or {hg.pct_loss_age_adh:.1f}% if routine services miss older children and adherence falls.
- **Up to about {h5.pct_children:.0f}% of children could move for a {h5.pct_loss:.0f}% loss** ({h5.pct_loss_age_adh:.0f}% with the age gap and adherence loss), all in {andj(h5.regions.split(', '))}.
- **The first districts to move would be** {', '.join(hyb_first.district.dropna().iloc[0:6].str.replace(chr(32) * 2, chr(32)))}: the natural candidates for a routine-delivery pilot.

![Hybrid delivery: children moved against cases averted lost](out/hybrid_curve.png)"""))

S.append(("equity", f"""## Who would lose protection

The door-to-door campaign reaches children regardless of their household. Routine delivery reaches a child only through routine services, which miss poorer children more often (first contact and monthly slips, from DHS 2022). Applied to {n_eq:,} northern children with dated vaccination cards (k = 2):

| Wealth quintile (national) | Share of children | Routine coverage per cycle (campaign 0.87) | Cases averted lost per 1,000 children | Share of all cases averted lost |
| --- | --- | --- | --- | --- |
{eq_rows}

- **Children in Ghana's poorest wealth quintile would bear {pct(eq2.loc[1, 'share_of_loss'])} of the lost protection** while making up {pct(eq2.loc[1, 'share_of_children'])} of these children; the richest quintile would lose almost nothing.
- **By mother's education the pattern is the same**: children of mothers with no schooling would bear {pct(eqe.loc['No education', 'share_of_loss'])} of the loss.
- **The concentration index of the loss is {eqs.loc[2.0, 'concentration_index']:.2f}** ({eqs.loc[2.8, 'concentration_index']:.2f} to {eqs.loc[1.4, 'concentration_index']:.2f} for k from 2.8 to 1.4): routine delivery would shift protection away from the poorest children, whom the campaign now reaches as well as anyone.

These children carry dated cards, so they are likely better off than those without; the true gradient may be steeper.

![Routine coverage and lost protection by household wealth](out/equity.png)"""))

S.append(("ml", f"""## Machine learning: a better map, the same answer

Two machine-learning analyses tested the spatial model and asked why routine contacts slip.

**A stacked model of routine-contact slips.** Random forest, gradient boosting and ridge-penalised logistic regression were trained on 16 satellite and access covariates (night-time lights, distance to roads and waterways, land cover, facility and CHPS density, terrain, rainfall). Their combined prediction was scored by the same leave-one-district-out test as the spatial model, with no district predicted by a model that had seen it.

| Model | Mean absolute error | Correlation | Log score |
| --- | --- | --- | --- |
{mlrows}

- {ml_claim}: closer in {int(mls.loc['ml_alone_stack', 'closer_than_main'])} of {int(mls.loc['ml_alone_stack', 'districts'])} districts (Wilcoxon p = {mls.loc['ml_alone_stack', 'wilcoxon_p_vs_main']:.2f}). Adding a spatial field to it does not help further.
- **The district map barely changes.** With the stacked model in step 4, {stk_cls.get('green', 0)} districts are green, {stk_cls.get('amber', 0)} amber and {stk_cls.get('red', 0)} red; {n_stk_same} of {n_d} keep their class (ratio correlation {ratio_corr:.2f}) and no never-green district becomes green. Changes: {stk_moves}. It is included as one of the {n_var} variants.

![Leave-one-district-out error: spatial models and machine learning](out/ml_stack_cv.png)

**What predicts a slipped contact.** A gradient-boosting model of {n_contacts:,} DHS contacts was explained with SHAP values. It discriminates no better than logistic regression (AUC {auc('gradient boosting'):.2f} vs {auc('logistic, all features'):.2f} in the north; travel time alone {auc('travel time only'):.2f}), so the relationships are close to linear.

- **Mothers and households explain the most**: mother {pct(slg['Mother'])} and household {pct(slg['Household'])} of the attribution in northern contacts, against {pct(slg['Place: access and density'])} for access and population density.
- **Mother's education and household wealth** lower the odds of a slip by a similar amount in every region (log-odds {min(eff('mother_edu', r) for r in ['northern','savannah','north east','upper east','upper west']):.2f} to {max(eff('mother_edu', r) for r in ['northern','savannah','north east','upper east','upper west']):.2f} for education, top vs bottom quarter).
- **Distance matters only at long walks**: walking time raises the odds by {eff('log2_tt_walk', 'northern'):.2f} in Northern but {eff('log2_tt_walk', 'upper east'):.2f} in Upper East, where CHPS are dense.
- **Region itself carries {pct(slr[(slr.feature == 'reg_northern') & (slr.region == 'northern')].share.iloc[0])} of the attribution in Northern**: contacts there slip more than household characteristics explain, which points to how services are run.

![What predicts a slipped routine contact (northern contacts, SHAP)](out/slip_shap.png)"""))

S.append(("check", f"""## What routine case data can and cannot confirm

District malaria data (DHIMS2) show the campaign working in children under 5, but they do not confirm the map's district ranking.

- **SMC flattens the under-5 season.** Once SMC starts, the share of a year's under-5 confirmed cases falling in June–October (the months SMC protects) should drop. In the old Northern region ({len(u5n)} districts; SMC from 2019) it {moved(u5n)} (2015–18 to 2020–24), falling in {int((u5n.change < 0).sum())} of {len(u5n)} districts. In Upper East (SMC from 2016; one earlier year, 2015) it {moved(u5ue)}, falling in {int((u5ue.change < 0).sum())} of {len(u5ue)}.
- **That drop does not track the map** (Spearman {val.loc['u5_oldnorthern','rho']:.2f}, p = {val.loc['u5_oldnorthern','p']:.2f}; within regions {val.loc['u5_oldnorthern_within','rho']:.2f}). The red districts show a clear drop (median {100*u5_cls.get('red', float('nan')):+.1f} points) and the amber ones none ({100*u5_cls.get('amber', float('nan')):+.1f}). That fits the map's message (the campaign works where routine delivery would fail), but it is not a test of the ranking.

Only under-5 data are used here. Year-to-year differences in rainfall timing cannot be separated from SMC with these data, and southern districts cannot serve as a comparison because they do not qualify for SMC (transmission there is perennial)."""))

S.append(("depends", f"""## What the answer depends on, and the data still needed

- **k, but only in part of the north.** At k = 1 (routine SMC as reliable as vaccination), {int(stress.loc[1.0,'green'])} districts would be green and none red; at k = 3, {int(stress.loc[3.0,'green'])} green and {int(stress.loc[3.0,'red'])} red. In Northern, Savannah and North East the break-even k is {rbe.loc['Northern','best']:.1f}–{rbe.loc['North East','best']:.1f}, below even the lowest evidence for k (1.4), so the answer there does not hinge on its exact value. Upper East and Upper West stay green only up to k of about {rbe.loc['Upper East','best']:.1f}–{rbe.loc['Upper West','best']:.1f}, against evidence near 2.
- **Reach of children aged 2–4.** Without it, no district reaches the 90% line.
- **Campaign coverage outside Upper East.** The measured 87% per cycle (Upper East) is applied elsewhere with a wide uncertainty range. At k = 2, routine delivery would match the campaign in Northern only if the campaign covered under {pct(rcc.loc['Northern', 'cc_k2_best'])} per cycle; district coverage data from NMEP would settle this.

Data that would firm this up:

- [x] DHIMS2 district-month malaria cases by age, all districts (received)
- [ ] NMEP: SMC coverage by district and cycle, 2022–2025
- [ ] DHIMS2: Penta and malaria vaccine doses by district and month, to estimate k by district
- [x] GHS facility list with coordinates, including CHPS (public facility finder; obtain GHS permission or the official DHIMS2 export before publication)
- [ ] A pilot of routine SMC (CHPS or child welfare clinics) in a few Upper East or Upper West districts, the only places where the decision depends on k"""))

S.append(("method", f"""## How the estimate was built

Five steps on a 1 km grid over the five SMC regions, summarised to {n_d} districts weighted by children aged 3–59 months. The work feeds the delivery axis of the SMC regimen-vs-delivery study.

1. **Access.** Least-cost travel time to the nearest public or faith-based facility on Malaria Atlas Project walking and motorised friction surfaces, using the Ghana Health Service public facility finder ({len(fac_n):,} facilities in the five regions, {int(fac_n.is_chps.sum()):,} of them CHPS; private facilities excluded). The earlier Maina et al. list (about 212 northern CHPS) is kept as a sensitivity variant.
2. **Routine contact timeliness.** Per-contact outcome from card-dated Penta doses 2 and 3 (DHS 2022): a contact slips if the gap exceeds 8 weeks or the dose was never given. Binomial INLA-SPDE model on log motorised travel time and population density, with a Matérn spatial field and a cluster effect; cluster covariates are under-5-weighted means over the DHS displacement buffer.
3. **Burden.** Observed under-5 confirmed malaria cases in June–October, the months SMC protects (DHIMS2, 2020–24 mean, by district) ÷ the under-5 population, applied to WorldPop 2025 children aged 3–59 months.
4. **Coverage and cases averted.** Campaign: 87% per cycle (Ayamba 2025), wider uncertainty outside Upper East. Routine: regional Penta1 × (1 − k × per-contact slip), k lognormal with median 2. Cases averted = no-SMC baseline × SPAQ protection 88% (Cairns 2021) × coverage.
5. **Classification.** Routine ÷ campaign cases averted per district, from 1,000 joint posterior draws; thresholds 90% and 70%; {n_var} variants for robustness; break-even k by district; a global sensitivity analysis (first-order variance shares, 20,000 draws); machine-learning checks (a stacked covariate model and SHAP analysis of slip predictors).

Within a district, burden and drug protection cancel from the ratio, so the class depends on the two coverages; burden sets the weighting and totals."""))

S.append(("estimates", f"""## Key estimates and checks

| Estimate or check | Result |
| --- | --- |
| Under-5s within 30 min walk of a CHPS | {', '.join(f'{r} {pct(chps30_reg[r])}' for r in NORTH)} |
| Per-contact slip, northern clusters | {pct(q_north)} overall; {pct(qb['<30'])} (<30 min walk) to {pct(qb['>120'])} (>2 h) |
| Odds ratio per doubling of motorised travel time (spatial model) | {ci(tt['or'], tt['lo'], tt['hi'])}; WAIC {tt['waic']:.0f} vs {tt['waic_no_field']:.0f} without the spatial field |
| Leave-one-district-out cross-validation ({int(dcv.loc['main','districts'])} northern districts): spatial model vs regional means | mean absolute error {dcv.loc['main','mae']:.3f} vs {dcv.loc['region','mae']:.3f}; correlation {dcv.loc['main','corr']:.2f} vs {dcv.loc['region','corr']:.2f}; log score {dcv.loc['main','logscore']:.3f} vs {dcv.loc['region','logscore']:.3f}; closer in {int(dcv.loc['main','closer_than_region'])} of {int(dcv.loc['main','districts'])} districts (p = {dcv.loc['main','wilcoxon_p_vs_region']:.2f}) |
| Spatial block cross-validation, northern clusters | correlation {cvn.loc['main_spatial','corr']:.2f}; log score {cvn.loc['main_spatial','logscore']:.3f} vs {cvn.loc['covariates','logscore']:.3f} covariates-only and {cvn.loc['intercept_only','logscore']:.3f} intercept-only |
| Routine per-cycle coverage (main, k uncertain) | {', '.join(f'{r} {rcov[r]:.2f}' for r in NORTH)}; all north {allcov:.2f} (campaign 0.87) |
| Contact due in the SMC season (June–September) vs other months, northern contacts | adjusted OR {ci(so['or'], so['lo'], so['hi'])}: no season effect; June {ssm.loc['Jun', 'or']:.2f}, September {ssm.loc['Sep', 'or']:.2f} |
| Rainfall in the month a dose fell due (CHIRPS) | OR {ci(rain.loc['rain100','or'], rain.loc['rain100','lo'], rain.loc['rain100','hi'])} per 100 mm: no effect |
| Rainy-season flooding (JRC surface water) | children gaining >30 min walking: {w_all.loc['main','pct_children_plus30']:.1f}% (main rule) to {w_all.loc['barrier','pct_children_plus30']:.1f}% (near-impassable bound) |
| Under-5 confirmed cases per child, June–October (DHIMS2 2020–24) | {', '.join(f'{r} {rbu.loc[r, "per_child"]:.2f}' for r in NORTH)} |
| Drop in under-5 June–October share after SMC vs map (old Northern, {len(u5n)} districts) | Spearman {val.loc['u5_oldnorthern','rho']:.2f} (p = {val.loc['u5_oldnorthern','p']:.2f}): no independent confirmation |

Distance and poverty cannot be fully separated: the map predicts where routine contacts run late, not that distance alone causes it."""))

S.append(("k", """## Evidence for k

Three comparisons put k near 2; one routine-data figure points lower. No study measures k for SMC itself.

| Comparison | Malaria product | Comparator | Implied k | Source |
| --- | --- | --- | --- | --- |
| Ghana household surveys, dropout dose 1→3 | RTS,S 76% → 66% (13.2% lost) | Penta 99% → 93% (6.1% lost), same areas | 2.2 | [Asante, Lancet 2024](https://doi.org/10.1016/S0140-6736(24)00004-7); [Welaga, Malar J 2026, Table 5](https://doi.org/10.1186/s12936-025-05778-9) |
| Ghana admin data Jan–Jun 2021, stand-alone later visit | RTS,S dose 3→4 ≈30% lost | Measles-rubella 87% → 77% (11.5% lost) | ≈2.6 | [WHO SAGE/MPAG evidence report, Sept 2021, p. 50–51](https://cdn.who.int/media/docs/default-source/immunization/mvip/full-evidence-report-on-the-rtss-as01-malaria-vaccine-for-sage-mpag-(sept2021).pdf) |
| Mali cluster RCT, SMC per round | Fixed-point 76.8–83.0% | Door-to-door 87.6–90.2% | 1.84 (miss ratio vs door-to-door; a check) | [Barry, PLoS One 2018](https://doi.org/10.1371/journal.pone.0193296) |
| Same WHO report, admin dropout dose 1→3 | RTS,S 1% | not reported | below 1 at face value | WHO report, p. 50 |

These compare dropout, while the model scales lateness. The Ghana figures were checked against the primary sources."""))

S.append(("sens", f"""## Sensitivity across {n_var} variants

![District classes under each model variant](out/summary_k_chart.png)

Every variant that changes an assumption towards realism (stricter lateness, the age gap, adherence, lower reach) removes green districts. Only the regional-means model, which uses Upper East and Upper West's own lower observed lateness, produces many green districts, all in those two regions."""))

S.append(("limits", f"""## Limitations

- **Proxy.** Routine SMC has never been measured in Ghana; its performance is inferred from vaccination timeliness scaled by k, which comes from malaria vaccine dropout.
- **Assumptions dominate the counts.** k, the age gap, adherence and campaign coverage outside Upper East move the class counts more than the spatial modelling does (the timeliness map explains {pct(G.loc[FQ].min())}–{pct(G.loc[FQ].max())} of the variance). The robust results are the geography and the break-even k, which do not depend on the value of k.
- **Causation.** Net of household wealth, distance has no clear independent effect on timeliness.
- **Facility coordinates.** The GHS facility finder is an undocumented public source with no stated licence, and some CHPS coordinates differ by more than 15 km from 2016 records. Switching from the earlier list (about 212 northern CHPS) to this one ({int(fac_n.is_chps.sum()):,}) changed the class of {n_fac_changed} of {n_d} districts.
- **Partial-course protection** has no sourced estimate (0.5 and 0 assumed).
- **Campaign coverage is measured only in Upper East.** The conclusion assumes campaigns elsewhere perform roughly as well; in Upper East and Upper West the break-even campaign coverage ({pct(rcc.loc['Upper East', 'cc_k2_best'])}–{pct(rcc.loc['Upper West', 'cc_k2_best'])} at k = 2) is close to the measured 87%.
- **No independent validation yet.** Routine under-5 case data show the campaign's effect but do not confirm the map's district ranking. Validation needs district SMC coverage (NMEP) or a routine-delivery pilot.
- **Burden** uses confirmed under-5 cases in public facilities, which understate true burden; this affects totals, not district classes.
- **Direct effect only**: transmission effects are left to the EMOD component.
- **Thresholds** of 90% and 70% are a policy choice."""))

S.append(("sources", """## Sources

- Ghana Statistical Service and ICF. [Ghana Demographic and Health Survey 2022](https://dhsprogram.com/) (child recode and GPS, under DHS Program authorisation)
- Ghana Health Service, DHIMS2. District-month confirmed malaria cases by age, 2015–2025, all 248 districts
- Malaria Atlas Project. [Friction surfaces 2019; *P. falciparum* incidence and prevalence 2025](https://malariaatlas.org/)
- WorldPop. [Age and sex structures, R2025A, 1 km](https://hub.worldpop.org/)
- Ghana Health Service. [Public facility finder](https://ghs.gov.gh/facility-finder), accessed 10 October 2026 (public and faith-based facilities, including CHPS)
- Maina J, et al. [Public health facilities in sub-Saharan Africa](https://data.humdata.org/dataset/health-facilities-in-sub-saharan-africa) (HDX), used for the sensitivity variant
- Climate Hazards Center. [CHIRPS v2.0 monthly rainfall](https://data.chc.ucsb.edu/products/CHIRPS-2.0/); EC JRC. [Global Surface Water v1.4](https://global-surface-water.appspot.com/)
- Ayamba EY, et al. [Evaluation of seasonal malaria chemoprevention implementation in the Upper East region of Northern Ghana](https://doi.org/10.1186/s12936-025-05322-9). Malar J 2025;24:123
- Cairns M, et al. [Effectiveness of SMC treatments when implemented at scale](https://doi.org/10.1371/journal.pmed.1003727). PLoS Med 2021;18:e1003727
- Huang S, Baker K, Ibinaiye T, et al. [Predictors of accessing seasonal malaria chemoprevention medicines through non-door-to-door distribution in Nigeria](https://doi.org/10.1186/s12936-024-04964-5). Malar J 2024;23:131
- Asante KP, et al. [Feasibility, safety, and impact of the RTS,S/AS01E malaria vaccine when implemented through national immunisation programmes](https://doi.org/10.1016/S0140-6736(24)00004-7). Lancet 2024;403:1660
- Welaga P, et al. [The MVIP study area in Ghana: results of a household survey prior to the introduction of the RTS,S/AS01 vaccine](https://doi.org/10.1186/s12936-025-05778-9). Malar J 2026
- WHO. [Full Evidence Report on the RTS,S/AS01 Malaria Vaccine, SAGE/MPAG, September 2021](https://cdn.who.int/media/docs/default-source/immunization/mvip/full-evidence-report-on-the-rtss-as01-malaria-vaccine-for-sage-mpag-(sept2021).pdf)
- Barry A, Issiaka D, Traore T, et al. [Optimal mode for delivery of seasonal malaria chemoprevention in Ouelessebougou, Mali: a cluster randomized trial](https://doi.org/10.1371/journal.pone.0193296). PLoS One 2018;13:e0193296"""))

# Typographic minus for negative numbers (e.g. "Spearman -0.20" -> "Spearman −0.20").
S = [(k, re.sub(r"(?<=[\s(])-(?=\d)", "−", t)) for k, t in S]

title = "Routine vs campaign SMC delivery in northern Ghana: results summary"
md = f"# {title}\n\nAbdulzeid Yen Anafo · {AS_OF}\n\n" + "\n\n".join(t for _, t in S) + "\n"
open(OUT_MD, "w").write(md)
pd.Series({k: t for k, t in S}).to_json("out/summary_sections.json")
print("wrote", OUT_MD)

# ---- Word rendering of the same Markdown ----------------------------------------------
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.2))
doc.styles["Normal"].font.size = Pt(11)
for lvl, size in ((1, 18), (2, 14)):
    st = doc.styles[f"Heading {lvl}"]; st.font.size = Pt(size); st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
for nm in ("Normal", "Heading 1", "Heading 2", "List Bullet", "List Number"):
    rpr = doc.styles[nm].element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(att), None)
    for att in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(att), "Calibri")

LINK = r"\[([^\]]+)\]\(((?:[^()]|\([^()]*\))+)\)"
TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\((?:[^()]|\([^()]*\))+\))")


def hyperlink(par, url, text):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "1F5FA8"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(t); h.append(r); par._p.append(h)


def rich(par, text, size=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2]); run.bold = True
        elif part.startswith("["):
            m = re.match(LINK, part); hyperlink(par, m.group(2), m.group(1)); continue
        elif part.startswith("*"):
            run = par.add_run(part[1:-1]); run.italic = True
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)
    return par


def table(rows):
    hdr, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]; c.text = ""; rich(c.paragraphs[0], f"**{h}**", size=9.5)
        s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), "E8EEF5")
        c._tc.get_or_add_tcPr().append(s)
    for r in body:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""; rich(cells[i].paragraphs[0], v, size=9.5)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "A6A6A6")):
            e.set(qn(f"w:{k}"), v)
        b.append(e)
    t._tbl.tblPr.append(b)
    doc.add_paragraph()


def split_row(line):
    """Split a pipe-table row, ignoring pipes inside link URLs (none expected) and keeping cell text."""
    return [c.strip() for c in line.strip().strip("|").split(" | ")]


lines = md.splitlines(); i = 0
while i < len(lines):
    ln = lines[i]
    if not ln.strip():
        i += 1; continue
    if ln.startswith("# "):
        doc.add_heading(ln[2:], level=1)
    elif ln.startswith("## "):
        doc.add_heading(ln[3:], level=2)
    elif ln.startswith("!["):
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", ln)
        doc.add_picture(m.group(2), width=Cm(15.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp = rich(doc.add_paragraph(), f"*{m.group(1)}.*", size=9); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|\s*-", lines[i]):
                rows.append(split_row(lines[i]))
            i += 1
        table(rows); continue
    elif re.match(r"^- \[( |x)\] ", ln):
        box = "☑ " if ln[3] == "x" else "☐ "
        rich(doc.add_paragraph(style="List Bullet"), box + ln[6:])
    elif ln.startswith("- "):
        rich(doc.add_paragraph(style="List Bullet"), ln[2:])
    elif re.match(r"^\d+\. ", ln):
        rich(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", ln))
    else:
        p = rich(doc.add_paragraph(), ln); p.paragraph_format.space_after = Pt(6)
    i += 1

doc.save(OUT_DOCX)
print("wrote", OUT_DOCX)
