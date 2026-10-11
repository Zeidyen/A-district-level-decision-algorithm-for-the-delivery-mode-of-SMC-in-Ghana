"""Word version of the manuscript Claude Doc.

Input: ../manuscript_export.md, the Markdown export of the doc (figures appear there only as
"[image: Figure N. caption]" placeholders and equations as ```latex blocks).
Output: ../manuscript.docx, with the figures taken from out/ and the equations rendered as images.
Run from geo/: python3 scripts/23_manuscript_docx.py
"""
import re
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC, OUT = "../manuscript_export.md", "../manuscript.docx"
BYLINE = "Abdulzeid Yen Anafo · 10 October 2026"
FIG = {"Figure 1": "out/delivery_algorithm_flow.png", "Figure 2": "out/access_map.png",
       "Figure 3": "out/timeliness_map.png", "Figure 4": "out/summary_district_map.png",
       "Figure 5": "out/breakeven_map.png", "Figure 6": "out/breakeven_cc.png", "Figure 7": "out/gsa_chart.png",
       "Figure 8": "out/equity.png", "Figure 9": "out/delivery_algorithm_map.png",
       "Figure S1": "out/summary_k_chart.png", "Figure S2": "out/district_cv.png", "Figure S3": "out/ml_stack_cv.png",
       "Figure S4": "out/slip_shap.png", "Figure S5": "out/slip_shap_by_region.png", "Figure S6": "out/season_slips.png",
       "Figure S7": "out/hybrid_curve.png", "Figure S8": "out/sobol_total.png"}
# matplotlib mathtext versions of the doc's two equations (no \operatorname or \left/\right).
EQ = [r"$\mathrm{logit}(q_i) = \beta_0 + \beta_1 \log_2 T_i + \beta_2 \log D_i + u_i + S(\mathbf{s}_i)\qquad(1)$",
      r"$C_r = r \times P_1 \times (1 - k\,q) \times [1 - s\,(1 - \rho)] \times m_{adh},\quad m_{adh} = \frac{a_r + (1 - a_r) f}{a_c + (1 - a_c) f}\qquad(2)$"]

md = open(SRC).read()
md = md.replace("&#91;", "[")
lines = md.splitlines()


def unescape(t):
    return re.sub(r"\\([\\`*_{}\[\]()#+\-.!|>~])", r"\1", t)


def eq_png(tex, i):
    fn = f"out/manuscript_eq{i}.png"
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, tex, fontsize=13)
    fig.savefig(fn, dpi=300, bbox_inches="tight", pad_inches=0.05, transparent=False)
    plt.close(fig)
    return fn


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.line_spacing = 1.15
for lvl, size in ((1, 16), (2, 13.5), (3, 11.5)):
    st = doc.styles[f"Heading {lvl}"]; st.font.size = Pt(size); st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
for nm in ("Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
    rpr = doc.styles[nm].element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(att), None)
    for att in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(att), "Calibri")

LINK = r"\[([^\]]+)\]\(((?:[^()]|\([^()]*\))+)\)"
TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\((?:[^()]|\([^()]*\))+\))")


def hyperlink(par, url, text):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "1F5FA8"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(t); h.append(r); par._p.append(h)


def rich(par, text, size=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(unescape(part[2:-2])); run.bold = True
        elif part.startswith("["):
            m = re.match(LINK, part); hyperlink(par, m.group(2), unescape(m.group(1))); continue
        elif part.startswith("*"):
            run = par.add_run(unescape(part[1:-1])); run.italic = True
        else:
            run = par.add_run(unescape(part))
        if size:
            run.font.size = Pt(size)
    return par


def split_row(line):
    cells, cur, depth = [], "", 0
    body = line.strip()[1:-1] if line.strip().endswith("|") else line.strip()[1:]
    for ch in body:                         # split on | outside link parentheses
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if ch == "|" and depth == 0 and not cur.endswith("\\"):
            cells.append(cur.strip()); cur = ""
        else:
            cur += ch
    cells.append(cur.strip())
    return cells


def table(rows):
    hdr, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]; c.text = ""; rich(c.paragraphs[0], f"**{h}**", size=9)
        s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), "E8EEF5")
        c._tc.get_or_add_tcPr().append(s)
    for r in body:
        cells = t.add_row().cells
        for i, v in enumerate(r[:len(hdr)]):
            cells[i].text = ""; rich(cells[i].paragraphs[0], v, size=9)
    doc.add_paragraph()


i = 0; neq = 0; seen_title = False
while i < len(lines):
    ln = lines[i]
    if not ln.strip():
        i += 1; continue
    if ln.startswith("```latex"):
        j = i + 1
        while not lines[j].startswith("```"):
            j += 1
        doc.add_picture(eq_png(EQ[neq], neq + 1), width=None)
        p = doc.paragraphs[-1]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic = p.runs[-1]._r.find(".//" + qn("wp:extent"))
        cx, cy = int(pic.get("cx")), int(pic.get("cy"))   # cap width at 14 cm
        if cx > Cm(14):
            sc = Cm(14) / cx; inl = p.runs[-1]._r.find(".//" + qn("wp:inline"))
            for e in (pic, inl.find(qn("wp:extent")), p.runs[-1]._r.find(".//" + qn("a:ext"))):
                if e is not None:
                    e.set("cx", str(int(cx * sc))); e.set("cy", str(int(cy * sc)))
        neq += 1; i = j + 1; continue
    if ln.startswith("# ") and not seen_title:
        doc.add_heading(unescape(ln[2:]), level=1); seen_title = True
        rich(doc.add_paragraph(), BYLINE)
        if i + 2 < len(lines) and "@" in lines[i + 2]:
            i += 2                                            # skip the doc's own byline
    elif ln.startswith("### "):
        doc.add_heading(unescape(ln[4:]), level=3)
    elif ln.startswith("## "):
        doc.add_heading(unescape(ln[3:]), level=2)
    elif ln.startswith("[image: "):
        cap = unescape(ln[len("[image: "):]).rstrip("]").rstrip()
        key = re.match(r"(Figure S?\d+)", cap).group(1)
        doc.add_picture(FIG[key], width=Cm(15.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
        p = doc.add_paragraph(); r = p.add_run(key + "."); r.bold = True; r.font.size = Pt(9)
        r2 = p.add_run(cap[len(key) + 1:]); r2.font.size = Pt(9)
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|\s*:?-", lines[i]):
                rows.append(split_row(lines[i]))
            i += 1
        table(rows); continue
    elif re.match(r"^\s*[-*] ", ln):
        rich(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", ln))
    elif re.match(r"^\d+\. ", ln):
        n, txt = re.match(r"^(\d+)\. (.*)", ln).groups()
        p = rich(doc.add_paragraph(), f"{n}. " + txt)           # keep reference numbers literal
        p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(2)
    else:
        p = rich(doc.add_paragraph(), ln); p.paragraph_format.space_after = Pt(6)
    i += 1

doc.save(OUT)
print("wrote", OUT, "| equations:", neq)
